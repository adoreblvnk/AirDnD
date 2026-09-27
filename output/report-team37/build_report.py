from pathlib import Path
import json, re, html, csv, hashlib
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.mathtext import math_to_image
from matplotlib.font_manager import FontProperties
from PIL import Image as PILImage
from matplotlib.patches import FancyBboxPatch
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER, TA_LEFT
from reportlab.platypus import Paragraph, Image, Table, TableStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Pt, Mm, Inches
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[1]
FIG=OUT/'figures'
REPORT=json.loads((ROOT/'evidence/reports/benchmark.json').read_text())
from report_content import TITLE, AUTHORS, COLUMNS
from draw_concepts import FIG

TABLE=[['Method','Leakage (%)','Duplicates / run']]
for m,n in zip(['naive_static','independent_greedy','deterministic_ablation','airdnd','ortools_teacher'],['Naive static','Independent greedy','Handwritten beliefs','AirDnD','OR-Tools reference']):
 TABLE.append([n,f"{REPORT['statistics'][m]['leakage']['mean']*100:.2f}",f"{REPORT['statistics'][m]['duplicate_pursuit']['mean']:.2f}"])
equation_paths={}
latex_equations=[]
for blocks in COLUMNS:
 for b in blocks:
  if b[0]=='eq':
   fig=plt.figure(figsize=(3.4,.48))
   fig.text(.02,.48,'$'+b[1]+'$',ha='left',va='center',fontsize=10)
   number=len(equation_paths)+1
   equation_paths[b[1]]=FIG/f'equation-{number}.png'
   fig.text(.98,.48,f'({number})',ha='right',va='center',fontsize=9)
   fig.savefig(equation_paths[b[1]],dpi=400,facecolor='white')
   fig.savefig(FIG/f'equation-{number}.pdf',facecolor='white')
   plt.close(fig)
   latex_equations.append('\\begin{equation}\n'+b[1]+'\n\\end{equation}\n')
(OUT/'formulas.tex').write_text('\n'.join(latex_equations),encoding='utf-8')

for n,f in [('TNR','times.ttf'),('TNR-Bold','timesbd.ttf'),('TNR-Italic','timesi.ttf'),('TNR-BoldItalic','timesbi.ttf')]:
 pdfmetrics.registerFont(TTFont(n,str(Path('C:/Windows/Fonts')/f)))
pdfmetrics.registerFontFamily('TNR',normal='TNR',bold='TNR-Bold',italic='TNR-Italic',boldItalic='TNR-BoldItalic')

def pdf_inline_math(text):
 def replace(match):
  latex=match.group(1)
  path=FIG/('inline-'+hashlib.sha256(latex.encode()).hexdigest()[:12]+'.png')
  math_to_image('$'+latex+'$',str(path),prop=FontProperties(size=10,math_fontfamily='stix'),dpi=400)
  with PILImage.open(path) as im: width,height=im.size
  return f'<img src="{path.as_posix()}" width="{width*72/400}" height="{height*72/400}" valign="-2"/>'
 return re.sub(r'\$([^$]+)\$',replace,text)

def word_math_run(text):
 run=OxmlElement('m:r');node=OxmlElement('m:t');node.text=text;run.append(node)
 return run

def word_inline_math(paragraph,text):
 for index,part in enumerate(re.split(r'\$([^$]+)\$',text)):
  if index%2==0:
   paragraph.add_run(part)
   continue
  math=OxmlElement('m:oMath')
  if part==r'\hat{p}_k':
   base=OxmlElement('m:acc');props=OxmlElement('m:accPr');char=OxmlElement('m:chr');char.set(qn('m:val'),'\u0302');props.append(char);base.append(props)
   expr=OxmlElement('m:e');expr.append(word_math_run('p'));base.append(expr)
  else:
   base=word_math_run(part.split('_')[0])
  if '_' in part:
   subscript=OxmlElement('m:sSub');expr=OxmlElement('m:e');expr.append(base);subscript.append(expr)
   sub=OxmlElement('m:sub');sub.append(word_math_run(part.split('_')[1]));subscript.append(sub);math.append(subscript)
  else:math.append(base)
  paragraph._p.append(math)
ST={
'p':ParagraphStyle('p',fontName='TNR',fontSize=10,leading=11.2,alignment=TA_JUSTIFY,firstLineIndent=10,spaceAfter=4),
'abstract':ParagraphStyle('abstract',fontName='TNR',fontSize=9,leading=10.4,alignment=TA_JUSTIFY,spaceAfter=6),
'keywords':ParagraphStyle('keywords',fontName='TNR',fontSize=9,leading=10.4,spaceAfter=7),
'h':ParagraphStyle('h',fontName='TNR',fontSize=10,leading=11.5,alignment=TA_CENTER,spaceBefore=7,spaceAfter=5),
'sub':ParagraphStyle('sub',fontName='TNR-Italic',fontSize=10,leading=11.2,spaceBefore=3,spaceAfter=4),
'caption':ParagraphStyle('caption',fontName='TNR',fontSize=8,leading=9.2,spaceAfter=7),
'ref':ParagraphStyle('ref',fontName='TNR',fontSize=8,leading=9.1,spaceAfter=4,leftIndent=12,firstLineIndent=-12,wordWrap='LTR'),
'note':ParagraphStyle('note',fontName='TNR-Italic',fontSize=8,leading=9.2,spaceBefore=3,spaceAfter=4),
}
W,H=A4;M=45.4;GAP=14.2;CW=(W-2*M-GAP)/2;BOTTOM=42.5;TOP=H-45.4

def table_flow():
 cells=[[Paragraph(html.escape(s),ParagraphStyle('t',fontName='TNR-Bold' if ri==0 else 'TNR',fontSize=8,leading=9)) for s in row] for ri,row in enumerate(TABLE)]
 t=Table(cells,colWidths=[CW*.42,CW*.27,CW*.31]);t.setStyle(TableStyle([
 ('LINEABOVE',(0,0),(-1,0),.6,colors.black),('LINEBELOW',(0,0),(-1,0),.4,colors.black),('LINEBELOW',(0,-1),(-1,-1),.6,colors.black),
 ('LEFTPADDING',(0,0),(-1,-1),2),('RIGHTPADDING',(0,0),(-1,-1),2),('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3),('VALIGN',(0,0),(-1,-1),'TOP')]))
 return t

pdf=canvas.Canvas(str(OUT/'Team37_AirDnD_IEEE_Report.pdf'),pagesize=A4)
pdf.setTitle(TITLE);pdf.setAuthor('Team 37 — '+AUTHORS)
layout=[]
for page in range(3):
 top=TOP
 if page==0:
  title=Paragraph(TITLE,ParagraphStyle('title',fontName='TNR',fontSize=24,leading=26,alignment=TA_CENTER))
  _,th=title.wrap(W-2*M,100);title.drawOn(pdf,M,top-th);top-=th+10
  pdf.setFont('TNR',11);pdf.drawCentredString(W/2,top,'Team 37');top-=15
  pdf.drawCentredString(W/2,top,AUTHORS);top-=14
  pdf.setFont('TNR',9);pdf.drawCentredString(W/2,top,'Singapore Defense Tech Hackathon 2026 · Track 3, Layer 03');top-=21
 for col in range(2):
  x=M+col*(CW+GAP);y=top
  for b in COLUMNS[page*2+col]:
   typ=b[0]
   if typ=='fig':
    im=Image(str(FIG/b[1]),width=CW,height=b[2]);im.drawOn(pdf,x,y-b[2]);y-=b[2]+3
    q=Paragraph(b[3],ST['caption']);_,hh=q.wrap(CW,1000);q.drawOn(pdf,x,y-hh);y-=hh+ST['caption'].spaceAfter
   elif typ=='eq':
    im=Image(str(equation_paths[b[1]]),width=CW,height=34.6);im.drawOn(pdf,x,y-34.6);y-=40
   elif typ=='table':
    q=Paragraph('TABLE I<br/>INITIAL SIMULATION COMPARISON',ParagraphStyle('tc',parent=ST['caption'],alignment=TA_CENTER));_,hh=q.wrap(CW,100);q.drawOn(pdf,x,y-hh);y-=hh+3
    q=table_flow();_,hh=q.wrap(CW,1000);q.drawOn(pdf,x,y-hh);y-=hh+7
   else:
    sty=ST[typ];y-=sty.spaceBefore;q=Paragraph(pdf_inline_math(b[1]),sty);_,hh=q.wrap(CW,1500);q.drawOn(pdf,x,y-hh);y-=hh+sty.spaceAfter
  layout.append({'page':page+1,'column':col+1,'bottom_y':round(y,2),'remaining_pt':round(y-BOTTOM,2)})
  if y<BOTTOM: print('OVERFLOW',layout[-1])
 pdf.setFont('TNR',8);pdf.drawCentredString(W/2,25,str(page+1));pdf.showPage()
pdf.save()

def plain(s):return html.unescape(re.sub('<[^>]+>','',s).replace('<br/>',' '))
doc=Document();sec=doc.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297);sec.top_margin=Pt(45.4);sec.bottom_margin=Pt(42.5);sec.left_margin=Pt(M);sec.right_margin=Pt(M)
sec.header_distance=Pt(16);sec.footer_distance=Pt(20)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(10)
normal.paragraph_format.line_spacing=Pt(11.2);normal.paragraph_format.space_after=Pt(4)
normal.paragraph_format.widow_control=False
for style in ['Title','Heading 1','Heading 2','Caption']:
 doc.styles[style].font.name='Times New Roman'
q=doc.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_after=Pt(10);q.paragraph_format.line_spacing=Pt(26)
r=q.add_run(TITLE);r.font.size=Pt(24)
for text,size,after in [('Team 37',11,3),(AUTHORS,11,3),('Singapore Defense Tech Hackathon 2026 · Track 3, Layer 03',9,10)]:
 q=doc.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_after=Pt(after);q.paragraph_format.line_spacing=Pt(size+3);r=q.add_run(text);r.font.size=Pt(size)
sec=doc.add_section(WD_SECTION_START.CONTINUOUS)
cols=sec._sectPr.find(qn('w:cols'));cols.set(qn('w:num'),'2');cols.set(qn('w:space'),str(round(GAP*20)))
for ci,blocks in enumerate(COLUMNS):
 if ci:
  q=doc.add_paragraph();q.paragraph_format.space_after=Pt(0);q.paragraph_format.line_spacing=Pt(1);q.add_run().add_break(WD_BREAK.PAGE if ci%2==0 else WD_BREAK.COLUMN)
 for b in blocks:
  typ=b[0]
  if typ=='fig':
   q=doc.add_paragraph();q.paragraph_format.space_after=Pt(3);q.paragraph_format.line_spacing=1
   q.add_run().add_picture(str(FIG/b[1]),width=Pt(CW),height=Pt(b[2]))
   q=doc.add_paragraph(b[3]);q.paragraph_format.line_spacing=Pt(9.2);q.paragraph_format.space_after=Pt(7)
   for r in q.runs:r.font.size=Pt(8)
  elif typ=='eq':
   q=doc.add_paragraph();q.paragraph_format.space_after=Pt(4);q.paragraph_format.line_spacing=1
   q.add_run().add_picture(str(equation_paths[b[1]]),width=Pt(CW),height=Pt(34.6))
  elif typ=='table':
   q=doc.add_paragraph('TABLE I\nINITIAL SIMULATION COMPARISON');q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.line_spacing=Pt(9.2)
   for r in q.runs:r.font.size=Pt(8)
   t=doc.add_table(rows=0,cols=3);t.autofit=False
   for ri,row in enumerate(TABLE):
    cells=t.add_row().cells
    for j,s in enumerate(row):
     cells[j].width=Pt(CW*[.42,.27,.31][j]);q=cells[j].paragraphs[0];q.paragraph_format.line_spacing=Pt(9);q.paragraph_format.space_after=Pt(3);r=q.add_run(s);r.font.size=Pt(8);r.bold=ri==0
   doc.add_paragraph().paragraph_format.space_after=Pt(2)
  else:
   q=doc.add_paragraph();word_inline_math(q,plain(b[1]));sty=ST[typ];q.paragraph_format.space_before=Pt(sty.spaceBefore);q.paragraph_format.space_after=Pt(sty.spaceAfter);q.paragraph_format.line_spacing=Pt(sty.leading)
   q.alignment=WD_ALIGN_PARAGRAPH.CENTER if typ=='h' else WD_ALIGN_PARAGRAPH.JUSTIFY if typ in ['p','abstract'] else WD_ALIGN_PARAGRAPH.LEFT
   if typ=='p':q.paragraph_format.first_line_indent=Pt(10)
   if typ=='ref':q.paragraph_format.left_indent=Pt(12);q.paragraph_format.first_line_indent=Pt(-12)
   for r in q.runs:r.font.size=Pt(sty.fontSize);r.italic=typ in ['sub','note'];r.bold=typ=='abstract'
footer=doc.sections[0].footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.append(field)
doc.core_properties.title=TITLE;doc.core_properties.author='Team 37 — '+AUTHORS
doc.save(OUT/'Team37_AirDnD_IEEE_Report.docx')

md=['# '+TITLE,'','Team 37 — '+AUTHORS,'']
for blocks in COLUMNS:
 for b in blocks:
  if b[0]=='fig':md.extend([f'![{b[3]}](figures/{b[1]})',''])
  elif b[0]=='eq':md.extend(['$$',b[1],'$$',''])
  elif b[0]=='table':
   md.extend(['| '+' | '.join(TABLE[0])+' |','|---|---:|---:|']+['| '+' | '.join(row)+' |' for row in TABLE[1:]]+[''])
  else:md.extend([('## ' if b[0]=='h' else '### ' if b[0]=='sub' else '')+plain(b[1]),''])
(OUT/'report.md').write_text('\n'.join(md),encoding='utf-8')
(OUT/'layout-check.json').write_text(json.dumps(layout,indent=2),encoding='utf-8')
print(json.dumps(layout,indent=2));print('Words:',len(' '.join(md).split()))
